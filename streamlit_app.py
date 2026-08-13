import os, re, json, io
from datetime import datetime

import streamlit as st
import pandas as pd
from google import genai
from PIL import Image
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Lieferzettel KI", page_icon="🚚", layout="wide")

MODEL = os.getenv("GEMINI_MODEL", "gemini-3.5-flash-lite")

COLUMNS = [
    "Pos.", "Name", "Straße", "PLZ", "Ort", "Bis", "Colli",
    "Gewicht", "Ablageplatz", "Schlüssel-Nr.", "Zeit"
]

PROMPT = r"""
Du liest eine deutsche Wegliste als Tabelle.

SEHR WICHTIG:
Die allererste Zeile der Tabelle ist eine EINZIGE ZUSAMMENHÄNGENDE KOPFZEILE über allen Spalten.
Sie gehört zur Tabelle und ist KEINE eigene Lieferung.

Danach folgen die Spaltenüberschriften:
Pos., Name, Straße, PLZ, Ort, Bis, Colli, Gew., Ablageplatz, Schlüssel-Nr., Zeit.

Danach kommen die Lieferzeilen.

Regeln:
1. Jede sichtbare Lieferzeile als eigenes Objekt ausgeben.
2. Keine sichtbare Lieferzeile überspringen.
3. Wenn eine Zelle nicht sicher lesbar ist, nur diese Zelle leer lassen.
4. Niemals Werte aus einer Nachbarzeile übernehmen.
5. Niemals Straße als Name verwenden oder umgekehrt.
6. Niemals raten oder ähnliche Zeilen automatisch ergänzen.
7. Eine Lieferzeile kann auch ohne lesbare Positionsnummer existieren. Dann pos="" und die übrigen sichtbaren Felder trotzdem erfassen.
8. Überschriften/Leerzeilen unterhalb der Lieferdaten sind keine Lieferungen.
9. Erfinde keine Informationen.

Gib ausschließlich gültiges JSON zurück:
{
  "kopfzeile": "",
  "rows": [
    {
      "pos": "",
      "name": "",
      "strasse": "",
      "plz": "",
      "ort": "",
      "bis": "",
      "colli": "",
      "gewicht": "",
      "ablageplatz": "",
      "schluessel_nr": "",
      "zeit": ""
    }
  ]
}
"""

def get_client():
    key = st.secrets.get("GEMINI_API_KEY", os.getenv("GEMINI_API_KEY", ""))
    if not key:
        st.error("Der Gemini-Key ist noch nicht als Secret hinterlegt.")
        st.stop()
    return genai.Client(api_key=key)

def extract(image):
    key = st.secrets.get("GEMINI_API_KEY", os.getenv("GEMINI_API_KEY", ""))

    if not key:
        st.error("Der Gemini-Key ist nicht als Secret hinterlegt.")
        st.stop()

    client = genai.Client(api_key=key)

    image = Image.open(image)

    response = client.models.generate_content(
        model=MODEL,
        contents=[PROMPT, image]
    )

    raw = response.text.strip()
    raw = re.sub(r"^```(?:json)?\s*", "", raw)
    raw = re.sub(r"\s*```$", "", raw)

    return json.loads(raw)

    return json.loads(raw)

def make_xlsx(header, df):
    wb = Workbook()
    ws = wb.active
    ws.title = "Wegliste"

    clean = df[COLUMNS].fillna("")
    last_col = len(COLUMNS)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(1, 1, str(header or "Wegliste"))
    ws.cell(1, 1).font = Font(bold=True)
    ws.cell(1, 1).alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 24

    for c, name in enumerate(COLUMNS, 1):
        cell = ws.cell(2, c, name)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r, row in enumerate(clean.itertuples(index=False), 3):
        for c, value in enumerate(row, 1):
            ws.cell(r, c, str(value) if value is not None else "")

    thin = Side(style="thin")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=last_col):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="center")

    widths = [8, 34, 28, 10, 22, 10, 10, 12, 16, 20, 10]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "A3"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()

st.title("🚚 Lieferzettel KI")
st.caption("Foto/Kamera → KI → Tabelle prüfen → korrigieren → echte Tabelle herunterladen")

uploaded = st.file_uploader("📁 Foto auswählen", type=["jpg", "jpeg", "png", "webp"])
camera = st.camera_input("📷 Oder direkt mit der Kamera fotografieren")

image = camera if camera is not None else uploaded

if image is not None:
    if st.button("🤖 Lieferzettel auslesen", type="primary"):
        with st.spinner("Lieferzettel wird gelesen …"):
            try:
                data = extract(image)
                st.session_state["header"] = str(data.get("kopfzeile", "") or "")
                rows = data.get("rows", [])
                mapping = {
                    "pos":"Pos.", "name":"Name", "strasse":"Straße", "plz":"PLZ",
                    "ort":"Ort", "bis":"Bis", "colli":"Colli", "gewicht":"Gewicht",
                    "ablageplatz":"Ablageplatz", "schluessel_nr":"Schlüssel-Nr.", "zeit":"Zeit"
                }
                records = []
                for row in rows:
                    records.append({col: str(row.get(key, "") or "") for key, col in mapping.items()})
                st.session_state["df"] = pd.DataFrame(records, columns=COLUMNS)
                st.success(f"✅ {len(records)} Lieferzeilen erkannt. Bitte kontrollieren.")
            except Exception as e:
                st.error(f"Fehler bei der Auslesung: {e}")

if "df" in st.session_state:
    st.subheader("📋 Wegliste")

    st.session_state["header"] = st.text_input(
        "Weglisten-Kopfzeile", value=st.session_state.get("header", "")
    )

    edited = st.data_editor(
        st.session_state["df"],
        num_rows="dynamic",
        use_container_width=True,
        key="delivery_editor"
    )
    st.session_state["df"] = edited

    st.info("✏️ Fehlende Buchstaben/Zahlen kannst du direkt in der Tabelle korrigieren.")

    xlsx = make_xlsx(st.session_state["header"], st.session_state["df"])
    st.download_button(
        "💾 Wegliste als echte Tabelle herunterladen",
        data=xlsx,
        file_name=f"Wegliste_{datetime.now().strftime('%Y-%m-%d_%H-%M-%S')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary"
    )
